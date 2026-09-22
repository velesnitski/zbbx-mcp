"""The sheet-writing helpers in excel.py, driven through a real openpyxl
workbook: what lands in which cell, with which style, and how wide the
columns end up. test_excel.py covers the pure classifiers.
"""

from __future__ import annotations

from openpyxl import Workbook

from zbbx_mcp.excel import auto_width, finalize_sheet, write_data_rows, write_headers

HEADERS = ["#", "Host", "Country", "Traffic In Mbps"]
ROWS = [
    {"Host": "srv-aq9001", "Country": "AQ", "Traffic In Mbps": 12.5},
    {"Host": "srv-bv9002", "Country": "BV"},
]


def _sheet():
    return Workbook().active


def _filled(rows=ROWS, headers=HEADERS):
    ws = _sheet()
    write_headers(ws, headers)
    write_data_rows(ws, rows, headers)
    return ws


class TestWriteHeaders:
    def test_header_row_carries_values_and_style(self):
        ws = _sheet()
        write_headers(ws, HEADERS)
        assert [ws.cell(1, c).value for c in range(1, len(HEADERS) + 1)] == HEADERS
        cell = ws.cell(1, 2)
        assert cell.font.bold is True
        assert cell.font.color.rgb.endswith("FFFFFF")
        assert cell.fill.fill_type == "solid"
        assert cell.fill.start_color.rgb.endswith("2F5496")
        assert cell.alignment.horizontal == "center"

    def test_no_headers_writes_nothing(self):
        ws = _sheet()
        write_headers(ws, [])
        assert ws.max_row == 1
        assert ws.cell(1, 1).value is None


class TestWriteDataRows:
    def test_rows_are_numbered_and_keyed_by_header(self):
        ws = _filled()
        # Column 1 is the running row number, never a value from the dict.
        assert ws.cell(2, 1).value == 1
        assert ws.cell(3, 1).value == 2
        assert ws.cell(2, 2).value == "srv-aq9001"
        assert ws.cell(2, 3).value == "AQ"
        assert ws.cell(2, 4).value == 12.5
        assert ws.cell(3, 2).value == "srv-bv9002"

    def test_missing_key_renders_as_empty_string(self):
        ws = _filled()
        assert ws.cell(3, 4).value == ""

    def test_data_cells_get_the_thin_bottom_border(self):
        ws = _filled()
        assert ws.cell(2, 2).border.bottom.style == "thin"
        assert ws.cell(3, 4).border.bottom.style == "thin"
        # The row-number column is written without a border.
        assert ws.cell(2, 1).border.bottom.style is None

    def test_no_rows_leaves_only_the_header(self):
        ws = _sheet()
        write_headers(ws, HEADERS)
        write_data_rows(ws, [], HEADERS)
        assert ws.max_row == 1

    def test_extra_keys_in_a_row_are_ignored(self):
        ws = _filled([{"Host": "srv-aq9001", "Unlisted": "x"}], ["#", "Host"])
        assert ws.max_column == 2
        assert ws.cell(2, 2).value == "srv-aq9001"


class TestAutoWidth:
    def test_width_is_the_longest_value_plus_padding(self):
        ws = _filled()
        auto_width(ws, HEADERS)
        assert ws.column_dimensions["B"].width == len("srv-aq9001") + 3
        # Header "Traffic In Mbps" is longer than any value beneath it.
        assert ws.column_dimensions["D"].width == len("Traffic In Mbps") + 3
        # "#" and the row numbers are one character each.
        assert ws.column_dimensions["A"].width == 1 + 3

    def test_width_is_capped(self):
        ws = _filled([{"Host": "x" * 80}], ["#", "Host"])
        auto_width(ws, ["#", "Host"])
        assert ws.column_dimensions["B"].width == 45

    def test_sample_rows_bounds_the_scan(self):
        rows = [{"Host": "a"}, {"Host": "b"}, {"Host": "c"}, {"Host": "y" * 30}]
        ws = _filled(rows, ["#", "Host"])
        auto_width(ws, ["#", "Host"], sample_rows=2)
        # Only the first two data rows are measured; the long one in row 5 is not.
        assert ws.column_dimensions["B"].width == len("Host") + 3
        auto_width(ws, ["#", "Host"])
        assert ws.column_dimensions["B"].width == 30 + 3

    def test_none_cells_do_not_count(self):
        ws = _filled([{"Host": None}], ["#", "Host"])
        auto_width(ws, ["#", "Host"])
        assert ws.column_dimensions["B"].width == len("Host") + 3

    def test_no_headers_is_a_noop(self):
        ws = _filled()
        auto_width(ws, [])
        assert len(ws.column_dimensions) == 0


class TestFinalizeSheet:
    def test_filter_freeze_and_widths_are_applied(self):
        ws = _filled()
        finalize_sheet(ws, HEADERS, len(ROWS))
        assert ws.auto_filter.ref == "A1:D3"
        assert ws.freeze_panes == "A2"
        assert ws.column_dimensions["B"].width == len("srv-aq9001") + 3

    def test_zero_rows_filters_the_header_only(self):
        ws = _sheet()
        write_headers(ws, ["#", "Host"])
        finalize_sheet(ws, ["#", "Host"], 0)
        assert ws.auto_filter.ref == "A1:B1"
        assert ws.freeze_panes == "A2"
