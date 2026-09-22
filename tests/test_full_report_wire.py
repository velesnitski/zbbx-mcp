"""``generate_full_report`` wire test: the fetch_all_data pipeline rendered as
an eight-sheet workbook. Fixtures are synthetic."""

from __future__ import annotations

import os
import time

import pytest
from openpyxl import load_workbook

from tests.wiretest import RecordingClient, run_tool
from zbbx_mcp import classify
from zbbx_mcp.tools import full_report

NOW = int(time.time())
LIVE = str(NOW)
STALE = str(NOW - 7200)

KEY_CPU = "system.cpu.util[,idle]"
KEY_LOAD = "system.cpu.load[percpu,avg5]"
KEY_MEM = "vm.memory.size[available]"
KEY_AGENT = "agent.version"
GB = 1_073_741_824


@pytest.fixture(autouse=True)
def _synthetic_env(monkeypatch, tmp_path):
    monkeypatch.setenv("ZABBIX_PRODUCT_MAP", "")
    monkeypatch.setattr(classify, "_PRODUCT_MAP", None)
    monkeypatch.setenv("ZABBIX_PROVIDER_CIDRS",
                       '{"Provider A": ["192.0.2.0/24"], "Provider B": ["198.51.100.0/24"]}')
    monkeypatch.setattr(classify, "_EXTRA_PROVIDER_NETS", None)
    monkeypatch.setenv("ZBBX_FILE_ROOTS", str(tmp_path))


def _host(hid, name, group, ip=None):
    return {"hostid": hid, "host": name, "name": name.upper(), "status": "0",
            "groups": [{"name": group}], "interfaces": [{"ip": ip}] if ip else []}


HOSTS = [
    _host("1", "srv-aq9001", "app_free", "192.0.2.1"),     # hot, saturated, on dashboard
    _host("2", "srv-aq9002", "app_free", "192.0.2.2"),     # quiet, stock-template NIC key
    _host("3", "srv-bv9001", "app_paid", "198.51.100.1"),  # agent silent, on dashboard
    _host("4", "srv-hm9001", "app_paid"),                  # no IP
    _host("5", "srv-tf9001", "Templates"),                 # unclassified: skipped
]

DASHBOARDS = [{
    "dashboardid": "10", "name": "Fleet Overview",
    "pages": [
        {"name": "Base", "widgets": [{"fields": [{"type": "6", "value": "g1"}]}]},
        {"name": "Bouvet", "widgets": [{"fields": [{"type": "6", "value": "g2"}]}]},
    ],
}]
GRAPHS = [
    {"graphid": "g1", "hosts": [{"hostid": "1"}]},
    {"graphid": "g2", "hosts": [{"hostid": "3"}]},
]


def _item(hid, value, clock=LIVE, key=None):
    it = {"hostid": hid, "lastvalue": str(value), "lastclock": clock}
    if key:
        it["key_"] = key
        it["itemid"] = f"{hid}-{key}"
    return it


METRICS = {
    KEY_CPU: [_item("1", 10), _item("2", 95), _item("3", 40, STALE), _item("4", 50)],
    KEY_LOAD: [_item("1", 3.0), _item("2", 0.1)],
    KEY_MEM: [_item("1", 2 * GB), _item("4", 4 * GB)],
    KEY_AGENT: [_item("1", "7.0.1"), _item("2", "7.0.1")],
}
TRAFFIC_IN = [
    _item("1", 700_000_000, key="net.if.in[eth0]"),
    _item("2", 300_000_000, key='net.if.in["eth0"]'),   # quoted form: missed by the fast filter
    _item("4", 550_000_000, key="net.if.in[eth0]"),
]
TRAFFIC_OUT = [
    _item("1", 100_000_000, key="net.if.out[eth0]"),
]
COST = [
    {"hostid": "1", "macro": "{$COST_MONTH}", "value": "100"},
    {"hostid": "3", "macro": "{$COST_MONTH}", "value": "50"},
    {"hostid": "1", "macro": "{$BW_LIMIT}", "value": "1000"},
]
TEMPLATES = [{"hostid": "1", "parentTemplates": [{"name": "Linux base"}]}]


def _client(hosts=HOSTS, dashboards=DASHBOARDS):
    def host_get(p):
        if "selectParentTemplates" in p:
            return [t for t in TEMPLATES if t["hostid"] in p["hostids"]]
        return hosts

    def item_get(p):
        hostids = set(p.get("hostids", []))
        search = p.get("search") or {}
        if search.get("key_") == "*net.if.in[*":
            return [t for t in TRAFFIC_IN if t["hostid"] in hostids]
        if search:
            return []
        key = (p.get("filter") or {}).get("key_")
        if isinstance(key, list):
            pool = TRAFFIC_IN if key[0].startswith("net.if.in[") else TRAFFIC_OUT
            return [t for t in pool if t["hostid"] in hostids and t["key_"] in key]
        return [i for i in METRICS.get(key, []) if i["hostid"] in hostids]

    return RecordingClient({
        "host.get": host_get, "dashboard.get": dashboards, "graph.get": GRAPHS,
        "item.get": item_get, "usermacro.get": COST,
    })


def _run(c, **kw):
    return run_tool(full_report, "generate_full_report", c, **kw)


def _path(out):
    return next(ln for ln in out.splitlines() if ln.startswith("**File:**")).split("`")[1]


def _sheet_rows(ws, key_col=0):
    header = [c.value for c in ws[1]]
    return header, {r[key_col]: dict(zip(header, r, strict=True)) for r in ws.iter_rows(min_row=2, values_only=True)}


class TestWire:
    def test_calls(self, tmp_path):
        c = _client()
        _run(c, output_dir=str(tmp_path))
        hosts = c.sent("host.get")
        assert hosts["filter"] == {"status": "0"} and hosts["selectGroups"] == ["name"]
        assert set(c.sent("usermacro.get")["filter"]["macro"]) == {"{$COST_MONTH}", "{$BW_LIMIT}"}
        assert c.sent("graph.get")["graphids"] and c.sent("graph.get")["selectHosts"] == ["hostid"]
        fallback = next(p for m, p in c.calls if m == "item.get" and (p.get("search") or {}).get("key_"))
        # only the hosts the fast key filter missed are re-read by wildcard
        assert "2" in fallback["hostids"] and "1" not in fallback["hostids"] and "4" not in fallback["hostids"]
        assert fallback["searchWildcardsEnabled"] is True


class TestHappyPath:
    def test_summary_text(self, tmp_path):
        out = _run(_client(), output_dir=str(tmp_path))
        assert "**Full Server Report**" in out
        assert "**Servers:** 4 distinct (2 on dashboards, 2 off)" in out
        assert "**Bandwidth:** 1 critical (>=650 Mbps), 1 high (>=500 Mbps)" in out
        assert "**Total cost:** $150.00/month" in out
        assert "**Countries:** 3 | **Products:** 2" in out
        assert "1. **All Servers** — 4 × 29 columns" in out
        assert "5. **Dashboard Tabs** — 2 tabs" in out
        assert "8. **Off-Dashboard** — 2 unmonitored servers" in out

    def test_workbook_sheets(self, tmp_path):
        out = _run(_client(), output_dir=str(tmp_path / "full"))
        path = _path(out)
        assert os.path.dirname(path) == os.path.realpath(tmp_path / "full")
        assert os.path.basename(path).startswith("zabbix_full_report_") and path.endswith(".xlsx")
        assert load_workbook(path).sheetnames == [
            "All Servers (4)", "Health Overview", "Product Analytics", "Country Analytics",
            "Dashboard Tabs", "Provider × Product", "Bandwidth Analysis", "Off-Dashboard (2)",
        ]

    def test_all_servers_sheet(self, tmp_path):
        ws = load_workbook(_path(_run(_client(), output_dir=str(tmp_path))))["All Servers (4)"]
        header, rows = _sheet_rows(ws, key_col=1)
        assert header[:3] == ["#", "Host", "Name"]
        assert list(rows) == ["srv-aq9001", "srv-bv9001", "srv-aq9002", "srv-hm9001"]   # dashboard hosts first
        r1 = rows["srv-aq9001"]
        assert (r1["Country"], r1["Dashboard"], r1["Tab"]) == ("AQ", "Fleet Overview", "Base")
        assert (r1["Product"], r1["Tier"], r1["Provider"], r1["IP"]) == ("app_free", "Default", "Provider A", "192.0.2.1")
        assert (r1["CPU %"], r1["Load Avg5"], r1["Mem Avail GB"]) == (90.0, 3.0, 2.0)
        assert (r1["Traffic In Mbps"], r1["Traffic Out Mbps"], r1["Traffic Total Mbps"]) == (700.0, 100.0, 800.0)
        assert (r1["BW Util %"], r1["BW Tier"]) == (70.0, "CRITICAL")     # against the host's own {$BW_LIMIT}
        assert (r1["Agent"], r1["Templates"]) == ("7.0.1", "Linux base")
        assert (r1["Cost/Month ($)"], r1["Cost/Year ($)"], r1["On Dashboard"]) == (100, 1200.0, "Yes")
        assert r1["All Tabs"] == "Fleet Overview / Base"
        r2 = rows["srv-aq9002"]
        assert (r2["Traffic In Mbps"], r2["BW Tier"], r2["On Dashboard"]) == (300.0, "NORMAL", "No")
        assert r2["BW Util %"] == 37.5                                     # default 800 Mbps ceiling
        r3 = rows["srv-bv9001"]
        assert r3["CPU %"] is None and r3["Traffic In Mbps"] is None and r3["BW Tier"] is None
        r4 = rows["srv-hm9001"]
        assert (r4["Provider"], r4["IP"], r4["BW Tier"]) == (None, None, "HIGH")

    def test_row_colours(self, tmp_path):
        ws = load_workbook(_path(_run(_client(), output_dir=str(tmp_path))))["All Servers (4)"]
        header = [c.value for c in ws[1]]
        cpu, tin = header.index("CPU %") + 1, header.index("Traffic In Mbps") + 1
        by_host = {ws.cell(row=r, column=2).value: r for r in range(2, 6)}
        assert ws.cell(row=by_host["srv-aq9001"], column=cpu).fill.start_color.rgb.endswith("FFC7CE")
        assert ws.cell(row=by_host["srv-aq9001"], column=tin).fill.start_color.rgb.endswith("FFC7CE")
        assert ws.cell(row=by_host["srv-aq9002"], column=cpu).fill.start_color.rgb.endswith("C6EFCE")
        assert ws.cell(row=by_host["srv-bv9001"], column=cpu).fill.fill_type is None

    def test_health_overview(self, tmp_path):
        ws = load_workbook(_path(_run(_client(), output_dir=str(tmp_path))))["Health Overview"]
        rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
        assert rows["Total Servers"][1] == 4
        assert rows["On Dashboard"][1] == 2 and rows["Off Dashboard"][1] == 2
        assert rows["CPU >= 80%"][1:3] == (1, "25.0%")
        assert rows["CPU < 10%"][1:3] == (1, "25.0%")
        assert rows["No CPU Data"][1:3] == (1, "25.0%")
        assert rows["Traffic >= 650 Mbps"][1] == 1 and rows["Traffic >= 500 Mbps"][1] == 2
        assert rows["No Traffic Data"][1] == 1
        assert rows["Unique Countries"][1] == 3 and rows["Unique Products"][1] == 2

    def test_product_analytics(self, tmp_path):
        ws = load_workbook(_path(_run(_client(), output_dir=str(tmp_path))))["Product Analytics"]
        header, rows = _sheet_rows(ws)
        free = rows["app_free"]
        assert (free["Servers"], free["On Dashboard"]) == (2, 1)
        assert (free["Median CPU %"], free["Max CPU %"]) == (47.5, 90.0)
        assert (free["Median Traffic Mbps"], free["Max Traffic Mbps"], free["Servers >= 650 Mbps"]) == (500.0, 700.0, 1)
        assert (free["Countries"], free["Providers"], free["Cost/Month ($)"]) == ("AQ", "Provider A", 100)
        paid = rows["app_paid"]
        assert (paid["Servers"], paid["Median CPU %"], paid["Servers >= 650 Mbps"]) == (2, 50.0, 0)
        assert rows["TOTAL"]["Servers"] == 4 and rows["TOTAL"]["Servers >= 650 Mbps"] == 1

    def test_country_analytics(self, tmp_path):
        ws = load_workbook(_path(_run(_client(), output_dir=str(tmp_path))))["Country Analytics"]
        header, rows = _sheet_rows(ws)
        assert list(rows)[0] == "AQ"                    # most servers first
        aq = rows["AQ"]
        assert (aq["Servers"], aq["Products"], aq["Providers"]) == (2, "app_free", "Provider A")
        assert (aq["Median CPU %"], aq["Total Traffic Gbps"], aq["Servers >= 650 Mbps"]) == (47.5, 1.0, 1)
        assert rows["BV"]["Median CPU %"] is None and rows["BV"]["Total Traffic Gbps"] is None
        assert rows["TOTAL"]["Servers"] == 4

    def test_dashboard_tabs(self, tmp_path):
        ws = load_workbook(_path(_run(_client(), output_dir=str(tmp_path))))["Dashboard Tabs"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert rows[0] == ("Fleet Overview", "Base", 1, 90.0, 700.0, 1, 0, 100)
        assert rows[1] == ("Fleet Overview", "Bouvet", 1, None, None, 0, 0, 50)
        assert ws.cell(row=2, column=6).fill.start_color.rgb.endswith("FFC7CE")

    def test_provider_product_matrix(self, tmp_path):
        ws = load_workbook(_path(_run(_client(), output_dir=str(tmp_path))))["Provider × Product"]
        header, rows = _sheet_rows(ws)
        assert header == ["Provider", "app_free/Default", "app_paid/Default", "Total Servers", "Total Cost ($)"]
        assert list(rows)[0] == "Provider A"
        assert rows["Provider A"]["app_free/Default"] == 2 and rows["Provider A"]["app_paid/Default"] is None
        assert (rows["Provider A"]["Total Servers"], rows["Provider A"]["Total Cost ($)"]) == (2, 100)
        assert (rows["No IP"]["app_paid/Default"], rows["No IP"]["Total Cost ($)"]) == (1, None)
        assert rows["Provider B"]["Total Cost ($)"] == 50

    def test_bandwidth_tiers(self, tmp_path):
        ws = load_workbook(_path(_run(_client(), output_dir=str(tmp_path))))["Bandwidth Analysis"]
        rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
        assert rows["CRITICAL"][2:] == (1, "25.0%", 90.0, 100)
        assert rows["HIGH"][2:4] == (1, "25.0%")
        assert rows["NORMAL"][2:4] == (1, "25.0%")
        assert rows["LOW"][2:4] == (0, "0.0%")
        assert rows["NO DATA"][2:] == (1, "25.0%", None, 50)

    def test_off_dashboard_sheet(self, tmp_path):
        ws = load_workbook(_path(_run(_client(), output_dir=str(tmp_path))))["Off-Dashboard (2)"]
        header, rows = _sheet_rows(ws, key_col=1)
        assert list(rows) == ["srv-aq9002", "srv-hm9001"]
        assert rows["srv-aq9002"]["Agent"] == "7.0.1" and rows["srv-aq9002"]["Traffic In Mbps"] == 300.0


class TestFilters:
    def test_product_filter(self, tmp_path):
        out = _run(_client(), product="app_paid", output_dir=str(tmp_path))
        assert "**Servers:** 2 distinct (1 on dashboards, 1 off)" in out
        assert "**Bandwidth:** 0 critical" in out
        assert "5. **Dashboard Tabs** — 1 tabs" in out
        assert "**Total cost:** $50.00/month" in out

    def test_country_filter(self, tmp_path):
        out = _run(_client(), country="bv", output_dir=str(tmp_path))
        assert "**Servers:** 1 distinct (1 on dashboards, 0 off)" in out
        assert "Off-Dashboard" not in load_workbook(_path(out)).sheetnames[-1]

    def test_include_and_exclude_lists(self, tmp_path):
        assert "**Servers:** 2 distinct" in _run(_client(), products="app_free", output_dir=str(tmp_path))
        assert "**Servers:** 2 distinct" in _run(_client(), exclude_product="app_free", output_dir=str(tmp_path))

    def test_off_dashboard_can_be_left_out(self, tmp_path):
        out = _run(_client(), include_off_dashboard=False, output_dir=str(tmp_path))
        assert "**Servers:** 2 distinct (2 on dashboards, 0 off)" in out
        assert "8. **Off-Dashboard** — 0 unmonitored servers" in out
        assert not any(s.startswith("Off-Dashboard") for s in load_workbook(_path(out)).sheetnames)


class TestEmptyFleet:
    def test_no_hosts(self, tmp_path):
        assert _run(_client(hosts=[], dashboards=[]), output_dir=str(tmp_path)) == "No servers found."

    def test_filter_that_matches_nothing(self, tmp_path):
        assert _run(_client(), country="pn", output_dir=str(tmp_path)) == "No servers found."

    def test_output_dir_outside_the_roots_is_refused(self):
        out = _run(_client(), output_dir="/nonexistent-root-for-test/x")
        assert out.startswith("Error:") and "not in the allowed roots" in out
