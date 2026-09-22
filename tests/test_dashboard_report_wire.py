"""``export_dashboard`` wire test: one dashboard's graphs resolve to hosts, the
hosts to a per-tab workbook. Fixtures are synthetic."""

from __future__ import annotations

import os
import time

import pytest
from openpyxl import load_workbook

from tests.wiretest import RecordingClient, run_tool
from zbbx_mcp import classify
from zbbx_mcp.tools import dashboard_report

NOW = int(time.time())
LIVE = str(NOW)
STALE = str(NOW - 7200)

KEY_CPU = "system.cpu.util[,idle]"
KEY_LOAD = "system.cpu.load[percpu,avg5]"
KEY_MEM = "vm.memory.size[available]"
GB = 1_073_741_824


@pytest.fixture(autouse=True)
def _synthetic_env(monkeypatch, tmp_path):
    monkeypatch.setenv("ZABBIX_PRODUCT_MAP", "")
    monkeypatch.setattr(classify, "_PRODUCT_MAP", None)
    monkeypatch.setenv("ZABBIX_PROVIDER_CIDRS",
                       '{"Provider A": ["192.0.2.0/24"], "Provider B": ["198.51.100.0/24"]}')
    monkeypatch.setattr(classify, "_EXTRA_PROVIDER_NETS", None)
    monkeypatch.setenv("ZBBX_FILE_ROOTS", str(tmp_path))


def _dashboard(name="Fleet Overview"):
    return [{
        "dashboardid": "77", "name": name,
        "pages": [
            {"name": "Base", "widgets": [
                {"fields": [{"type": "6", "value": "g1"}, {"type": "6", "value": "g2"}]},
                {"fields": [{"type": "6", "value": "g1"}]},           # same host twice: one row
                {"fields": [{"type": "6", "value": "g4"}]},           # graph with no host
            ]},
            {"name": "Bouvet", "widgets": [{"fields": [{"type": "6", "value": "g3"}]}]},
            {"name": "Empty", "widgets": []},
        ],
    }]


GRAPHS = [
    {"graphid": "g1", "hosts": [{"hostid": "1"}]},
    {"graphid": "g2", "hosts": [{"hostid": "2"}]},
    {"graphid": "g3", "hosts": [{"hostid": "3"}]},
    {"graphid": "g4", "hosts": []},
]


def _host(hid, name, group, ip):
    return {"hostid": hid, "host": name, "name": name.upper(), "status": "0",
            "groups": [{"name": group}], "interfaces": [{"ip": "127.0.0.1"}, {"ip": ip}]}


HOSTS = [
    _host("1", "srv-aq9001", "app_free", "192.0.2.1"),
    _host("2", "srv-aq9002", "app_free", "192.0.2.2"),
    _host("3", "srv-bv9001", "app_paid", "198.51.100.1"),
]


def _item(hid, value, clock=LIVE, key=None):
    it = {"hostid": hid, "lastvalue": str(value), "lastclock": clock}
    if key:
        it["key_"] = key
        it["itemid"] = f"{hid}-{key}"
    return it


METRICS = {
    KEY_CPU: [_item("1", 15), _item("2", 60), _item("3", 50, STALE)],
    KEY_LOAD: [_item("1", 2.0), _item("2", 0.75)],
    KEY_MEM: [_item("1", 2 * GB), _item("3", 6 * GB)],
}
TRAFFIC = [
    _item("1", 100_000_000, key="net.if.in[eth0]"),
    _item("1", 900_000_000, key="net.if.in[docker0]"),      # virtual NIC: must not count
    _item("2", 500_000_000, STALE, key="net.if.in[eth0]"),  # no longer reporting
    _item("3", 250_000_000, key="net.if.in[eth0]"),
]
COST = [
    {"hostid": "1", "value": "100"},
    {"hostid": "2", "value": "n/a"},     # unparsable macro is ignored
    {"hostid": "3", "value": "50"},
]


def _client(dashboards=None, graphs=GRAPHS):
    def item_get(p):
        if (p.get("search") or {}).get("key_") == "*net.if.in[*":
            return [t for t in TRAFFIC if t["hostid"] in p["hostids"]]
        key = (p.get("filter") or {}).get("key_")
        return [i for i in METRICS.get(key, []) if i["hostid"] in p["hostids"]]

    def host_get(p):
        return [h for h in HOSTS if h["hostid"] in p["hostids"]]

    return RecordingClient({
        "dashboard.get": _dashboard() if dashboards is None else dashboards,
        "graph.get": graphs, "host.get": host_get, "item.get": item_get, "usermacro.get": COST,
    })


def _run(c, **kw):
    return run_tool(dashboard_report, "export_dashboard", c, **kw)


def _path(out):
    return next(ln for ln in out.splitlines() if ln.startswith("**File:**")).split("`")[1]


class TestWire:
    def test_calls(self, tmp_path):
        c = _client()
        _run(c, dashboard_id="77", output_dir=str(tmp_path))
        assert c.sent("dashboard.get")["dashboardids"] == ["77"]
        assert c.sent("dashboard.get")["selectPages"] == "extend"
        assert set(c.sent("graph.get")["graphids"]) == {"g1", "g2", "g3", "g4"}
        assert set(c.sent("host.get")["hostids"]) == {"1", "2", "3"}
        traffic = next(p for m, p in c.calls if m == "item.get" and "search" in p)
        assert traffic["search"] == {"key_": "*net.if.in[*"} and traffic["searchWildcardsEnabled"] is True
        assert "lastclock" in traffic["output"]
        assert c.sent("usermacro.get")["filter"] == {"macro": "{$COST_MONTH}"}
        keys = [p["filter"]["key_"] for m, p in c.calls if m == "item.get" and "filter" in p and "search" not in p]
        assert sorted(keys) == sorted([KEY_CPU, KEY_LOAD, KEY_MEM])


class TestHappyPath:
    def test_summary_text(self, tmp_path):
        out = _run(_client(), dashboard_id="77", output_dir=str(tmp_path))
        assert "**Dashboard Export: Fleet Overview**" in out
        assert "**Servers:** 3 across 3 tabs" in out
        assert "- **Base**: 2 servers, median CPU 62.5%, median traffic 100.0 Mbps" in out
        assert "- **Bouvet**: 1 servers, median CPU N/A, median traffic 250.0 Mbps" in out
        assert "- **Empty**: 0 servers, median CPU N/A, median traffic N/A" in out
        assert "1. **All** — 3 servers × 14 columns" in out
        assert "2. **Base** — 2 servers" in out
        assert "3. **Bouvet** — 1 servers" in out
        assert "5. **Summary** — per-tab aggregates" in out

    def test_workbook_layout(self, tmp_path):
        out = _run(_client(), dashboard_id="77", output_dir=str(tmp_path / "x"))
        path = _path(out)
        assert os.path.dirname(path) == os.path.realpath(tmp_path / "x")
        assert os.path.basename(path).startswith("Fleet Overview_") and path.endswith(".xlsx")
        wb = load_workbook(path)
        assert wb.sheetnames == ["All (3)", "Base (2)", "Bouvet (1)", "Summary"]   # no sheet for an empty tab

    def test_all_sheet_rows(self, tmp_path):
        ws = load_workbook(_path(_run(_client(), dashboard_id="77", output_dir=str(tmp_path))))["All (3)"]
        header = [c.value for c in ws[1]]
        rows = {r[1]: dict(zip(header, r, strict=True)) for r in ws.iter_rows(min_row=2, values_only=True)}
        assert [r["#"] for r in rows.values()] == [1, 2, 3]
        r1 = rows["srv-aq9001"]
        assert (r1["Product"], r1["Tier"], r1["Provider"], r1["IP"]) == ("app_free", "Default", "Provider A", "192.0.2.1")
        assert (r1["CPU %"], r1["Load Avg5"], r1["Mem Avail GB"]) == (85.0, 2.0, 2.0)
        assert r1["Traffic Mbps"] == 100.0                     # the carrier NIC, not the virtual one
        assert r1["Cost/Month ($)"] == 100
        r2 = rows["srv-aq9002"]
        assert r2["Traffic Mbps"] is None                      # a stale rate is not traffic
        assert r2["Cost/Month ($)"] is None
        r3 = rows["srv-bv9001"]
        assert r3["CPU %"] is None and r3["Traffic Mbps"] == 250.0 and r3["Cost/Month ($)"] == 50
        assert ws.auto_filter.ref == "A1:N4" and ws.freeze_panes == "A2"

    def test_cpu_colouring(self, tmp_path):
        ws = load_workbook(_path(_run(_client(), dashboard_id="77", output_dir=str(tmp_path))))["All (3)"]
        cpu_col = [c.value for c in ws[1]].index("CPU %") + 1
        fills = {ws.cell(row=r, column=2).value: ws.cell(row=r, column=cpu_col).fill for r in range(2, 5)}
        assert fills["srv-aq9001"].start_color.rgb.endswith("FFC7CE")    # 85 -> red
        assert fills["srv-aq9002"].fill_type is None                     # 40: no colour band
        assert fills["srv-bv9001"].fill_type is None                     # no reading: no colour

    def test_summary_sheet(self, tmp_path):
        ws = load_workbook(_path(_run(_client(), dashboard_id="77", output_dir=str(tmp_path))))["Summary"]
        rows = {r[0]: r for r in ws.iter_rows(min_row=2, values_only=True)}
        assert rows["Base"] == ("Base", 2, 62.5, 100.0, 0, 100)
        assert rows["Bouvet"] == ("Bouvet", 1, None, 250.0, 0, 50)
        assert rows["Empty"] == ("Empty", 0, None, None, 0, None)
        assert rows["TOTAL"][:2] == ("TOTAL", 3)

    def test_tab_sheet_holds_only_its_hosts(self, tmp_path):
        wb = load_workbook(_path(_run(_client(), dashboard_id="77", output_dir=str(tmp_path))))
        assert [r[1] for r in wb["Base (2)"].iter_rows(min_row=2, values_only=True)] == ["srv-aq9001", "srv-aq9002"]
        assert [r[1] for r in wb["Bouvet (1)"].iter_rows(min_row=2, values_only=True)] == ["srv-bv9001"]

    def test_slashes_in_the_dashboard_name_cannot_escape_the_directory(self, tmp_path):
        out = _run(_client(dashboards=_dashboard("Fleet/Over\\view")), dashboard_id="77", output_dir=str(tmp_path))
        path = _path(out)
        assert os.path.dirname(path) == os.path.realpath(tmp_path)
        assert os.path.basename(path).startswith("Fleet-Over-view_")


class TestDegraded:
    def test_unknown_dashboard(self, tmp_path):
        assert _run(_client(dashboards=[]), dashboard_id="77", output_dir=str(tmp_path)) == "Dashboard '77' not found."

    def test_dashboard_without_graph_widgets(self, tmp_path):
        d = [{"dashboardid": "77", "name": "Fleet Overview",
              "pages": [{"name": "Base", "widgets": [{"fields": [{"type": "3", "value": "1"}]}]}]}]
        c = _client(dashboards=d)
        assert _run(c, dashboard_id="77", output_dir=str(tmp_path)) == "Dashboard 'Fleet Overview' has no graph widgets."
        assert "graph.get" not in {m for m, _ in c.calls}

    def test_graphs_that_resolve_to_no_host(self, tmp_path):
        c = _client(graphs=[{"graphid": "g1", "hosts": []}])
        assert _run(c, dashboard_id="77", output_dir=str(tmp_path)) == "No hosts resolved from dashboard graphs."
        assert "host.get" not in {m for m, _ in c.calls}

    def test_output_dir_outside_the_roots_is_refused(self):
        out = _run(_client(), dashboard_id="77", output_dir="/nonexistent-root-for-test/x")
        assert out.startswith("Error:") and "not in the allowed roots" in out
