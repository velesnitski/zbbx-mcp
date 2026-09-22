"""Wire tests for the cost audit tools (``tools/costs_audit.py``).

The audit tools read ``{$COST_MONTH}`` macros back and sort them into review
buckets or files. Each test pins the bucket a fixture row lands in, the
number computed from it, and the file written — an audit that files a row
under the wrong action is worse than none. Fixtures are synthetic.
"""

from __future__ import annotations

import csv
import json

import openpyxl
import pytest

from tests.wiretest import RecordingClient, run_tool
from zbbx_mcp import classify
from zbbx_mcp.tools import costs_audit

COST = "{$COST_MONTH}"


@pytest.fixture(autouse=True)
def _confined(monkeypatch, tmp_path):
    monkeypatch.setattr(classify, "_PRODUCT_MAP", {})
    monkeypatch.setenv("ZBBX_FILE_ROOTS", str(tmp_path))


def _host(hid, name, ip=None, group="app_free"):
    h = {"hostid": hid, "host": name, "name": name, "groups": [{"name": group}], "interfaces": []}
    if ip:
        h["interfaces"] = [{"ip": ip, "main": "1", "type": "1"}]
    return h


def _macro(hid, value, description=""):
    return {"hostmacroid": f"m{hid}", "hostid": hid, "macro": COST,
            "value": value, "description": description}


# Two documentation /24s carry hosts; the third (203.0.113.0/24) is empty so
# an address there has no neighbour signal.
HOSTS = [
    _host("1", "srv-aq9001", "192.0.2.10"),
    _host("2", "srv-aq9002", "192.0.2.11"),
    _host("3", "srv-bv9001", "198.51.100.20"),
    _host("4", "srv-hm9001", "198.51.100.30", group="app_paid"),
    _host("5", "mon-9001"),
]


class TestExportCostAudit:
    MACROS = [
        _macro("1", "20", "src:billing_ip"),         # billing-backed
        _macro("2", "25", "src:product_median"),     # estimated
        _macro("3", "210.82", ""),                   # legacy bulk value: estimated by heuristic
        _macro("4", "abc"),                          # unparsable: skipped
        _macro("99", "50"),                          # no such host: skipped
        _macro("5", "0"),                            # zero: skipped
    ]

    def _client(self, **overrides):
        return RecordingClient({"host.get": HOSTS, "usermacro.get": self.MACROS, **overrides})

    def test_estimated_mode_lists_only_unbacked_costs(self, tmp_path):
        out_xlsx = tmp_path / "audit.xlsx"
        out = run_tool(costs_audit, "export_cost_audit", self._client(), output_xlsx=str(out_xlsx))
        assert out == f"Exported 2 hosts (mode=estimated), $235.82/mo\nWrote: {out_xlsx}", out
        ws = openpyxl.load_workbook(out_xlsx).active
        assert ws.title == "cost_audit"
        assert [c.value for c in ws[1]] == ["Host", "Host ID", "IP", "Provider", "Product", "Tier",
                                            "Country", "Cost $/mo", "Source", "Billing-backed"]
        assert [c.value for c in ws[2]] == ["srv-bv9001", "3", "198.51.100.20", "Other", "app_free",
                                            "Default", "BV", 210.82, "(no tag)", "no"]
        assert [c.value for c in ws[3]] == ["srv-aq9002", "2", "192.0.2.11", "Other", "app_free",
                                            "Default", "AQ", 25, "src:product_median", "no"]
        assert ws.cell(row=5, column=1).value == "TOTAL (2 hosts)"
        assert ws.cell(row=5, column=8).value == 235.82

    def test_all_mode_includes_backed_costs(self, tmp_path):
        out_xlsx = tmp_path / "audit.xlsx"
        out = run_tool(costs_audit, "export_cost_audit", self._client(),
                       output_xlsx=str(out_xlsx), mode="all")
        assert out.startswith("Exported 3 hosts (mode=all), $255.82/mo"), out
        ws = openpyxl.load_workbook(out_xlsx).active
        assert [ws.cell(row=r, column=1).value for r in (2, 3, 4)] == ["srv-bv9001", "srv-aq9002", "srv-aq9001"]
        assert ws.cell(row=4, column=10).value == "yes"

    def test_source_of_truth_workbook_marks_and_links_rows(self, tmp_path):
        src = tmp_path / "source.xlsx"
        wb = openpyxl.Workbook()
        wb.active.title = "bill"
        wb.active.append(["line", "detail"])
        wb.active.append(["a", "192.0.2.11 / srv-aq9002"])      # host 2 is in the source
        wb.active.append(["b", "198.51.100.99"])                 # nobody
        wb.save(src)
        out_xlsx = tmp_path / "audit.xlsx"
        out = run_tool(costs_audit, "export_cost_audit", self._client(),
                       output_xlsx=str(out_xlsx), source_xlsx=str(src))
        assert "In source of truth: 1 yes (green) / 1 no — review needed" in out, out
        ws = openpyxl.load_workbook(out_xlsx).active
        assert ws.cell(row=1, column=11).value == "In source of truth"
        assert ws.cell(row=1, column=12).value == "Source row"
        # Row 2 is the 210.82 host (not in the source), row 3 the 25 host (in it).
        assert (ws.cell(row=2, column=11).value, ws.cell(row=2, column=12).value) == ("no", None)
        assert (ws.cell(row=3, column=11).value, ws.cell(row=3, column=12).value) == ("yes", "bill!2")
        assert ws.cell(row=3, column=12).hyperlink.target == f"file://{src.resolve()}#'bill'!A2"
        assert ws.cell(row=3, column=1).fill.fgColor.rgb.endswith("D9EAD3")
        assert ws.cell(row=2, column=1).fill.fill_type is None

    def test_failures_are_named(self, tmp_path):
        c = self._client()
        assert run_tool(costs_audit, "export_cost_audit", c, output_xlsx=str(tmp_path / "a.xlsx"),
                        source_xlsx=str(tmp_path / "missing.xlsx")).startswith("Failed to load source_xlsx:")
        assert run_tool(costs_audit, "export_cost_audit", c,
                        output_xlsx="/nonexistent-root/a.xlsx").startswith("Failed to write XLSX:")

        def boom(_p):
            raise ValueError("api down")
        assert run_tool(costs_audit, "export_cost_audit", self._client(**{"host.get": boom}),
                        output_xlsx=str(tmp_path / "a.xlsx")) == "Error: api down"


class TestDetectCostAnomalies:
    HOSTS = [*HOSTS, _host("6", "srv-tf9001", "198.51.100.40")]
    MACROS = [
        _macro("1", "20", "src:billing_ip"),
        _macro("2", "20"),
        _macro("3", "22"),
        _macro("4", "100", "src:billing_ip"),     # 5× the median
        _macro("5", "5"),                          # no address: not comparable
        _macro("6", "2", "src:bulk_pattern"),      # 0.1× the median
    ]

    def _client(self):
        return RecordingClient({"host.get": self.HOSTS, "usermacro.get": self.MACROS})

    def test_outliers_against_the_provider_median(self):
        c = self._client()
        out = run_tool(costs_audit, "detect_cost_anomalies", c)
        assert out.startswith("**2 cost anomalies detected**\nThresholds: >2.5× or <0.3× provider median"), out
        rows = [ln for ln in out.splitlines() if ln.startswith("| srv-")]
        assert rows == [
            "| srv-hm9001 | Other | $100.00 | $20.00 | 5.0× | src:billing_ip |",
            "| srv-tf9001 | Other | $2.00 | $20.00 | 0.1× | src:bulk_pattern |",
        ]
        assert [m for m, _ in c.calls] == ["host.get", "usermacro.get", "host.get", "usermacro.get"]

    def test_cap_and_quiet_fleet(self):
        out = run_tool(costs_audit, "detect_cost_anomalies", self._client(), max_results=1)
        assert "srv-tf9001" not in out and "*+1 more*" in out, out
        assert run_tool(costs_audit, "detect_cost_anomalies", self._client(),
                        high_factor=10, low_factor=0.01) == "No cost anomalies detected."


class TestAnalyzeCostImport:
    FILE = {
        "192.0.2.10": 30,                                             # already in Zabbix
        "192.0.2.77": {"name": "srv aq9002 dedicated", "price": 40},  # subnet + name: HIGH
        "198.51.100.99": 12,                                          # subnet only: MEDIUM
        "203.0.113.200": {"name": "unrelated", "price": 8},           # nothing: UNKNOWN
        "203.0.113.201": {"name": "aq9002 custom", "price": 5},       # name only: LOW
        "203.0.113.202": {"name": "xx aq9002", "price": 3},           # partial token: LOW
        "10.0.0.5": 50,                                               # private: ignored
        "203.0.113.203": "abc",                                       # unparsable: ignored
        "203.0.113.204": 0,                                           # free: ignored
    }

    def _run(self, tmp_path, data, hosts=HOSTS):
        src = tmp_path / "costs.json"
        src.write_text(json.dumps(data))
        out_json, out_csv = tmp_path / "a.json", tmp_path / "a.csv"
        c = RecordingClient({"host.get": hosts})
        out = run_tool(costs_audit, "analyze_cost_import", c, file_path=str(src),
                       output_csv=str(out_csv), output_json=str(out_json))
        return out, out_json, out_csv, c

    def test_tiers_scores_and_files(self, tmp_path):
        out, out_json, out_csv, c = self._run(tmp_path, self.FILE)
        assert out.startswith("**Cost Import Analysis: 5 unmatched IPs = $68.00/mo**"), out
        assert "| HIGH | 1 | $40.00 | $480.00 |" in out
        assert "| MEDIUM | 1 | $12.00 | $144.00 |" in out
        assert "| LOW | 2 | $8.00 | $96.00 |" in out
        assert "| UNKNOWN | 1 | $8.00 | $96.00 |" in out
        assert f"Saved: `{out_json}`" in out and f"Saved: `{out_csv}`" in out
        assert c.sent("host.get") == {"output": ["hostid", "host", "name"], "selectInterfaces": ["ip"]}

        results = json.loads(out_json.read_text())
        assert [(r["ip"], r["tier"], r["confidence"]) for r in results] == [
            ("192.0.2.77", "HIGH", 75), ("198.51.100.99", "MEDIUM", 40),
            ("203.0.113.201", "LOW", 35), ("203.0.113.202", "LOW", 15),
            ("203.0.113.200", "UNKNOWN", 0),
        ]
        assert results[0]["signals"] == ["/24 match: srv-aq9001 (+1 more)", "name match: srv-aq9002"]
        assert results[3]["signals"] == ["partial name: srv-aq9002"]
        assert results[4]["suggestion"].startswith("No signal")

        with open(out_csv, newline="") as f:
            rows = list(csv.reader(f))
        assert rows[0][:4] == ["Tier", "Confidence", "IP", "Billing Name"]
        assert rows[1][:6] == ["HIGH", "75", "192.0.2.77", "srv aq9002 dedicated", "40.00", "480.00"]
        assert len(rows) == 6

    def test_nothing_left_to_analyse(self, tmp_path):
        out, *_ = self._run(tmp_path, {"192.0.2.10": 0, "203.0.113.1": "x"})
        assert out == "No valid cost entries found in file."
        out, *_ = self._run(tmp_path, {"192.0.2.10": 30, "10.0.0.5": 9})
        assert out.startswith("**Cost Import Analysis: 0 unmatched IPs = $0.00/mo**"), out
        assert "| HIGH | 0 | $0.00 | $0.00 |" in out

    def test_missing_file(self, tmp_path):
        c = RecordingClient({"host.get": HOSTS})
        out = run_tool(costs_audit, "analyze_cost_import", c, file_path=str(tmp_path / "nope.json"),
                       output_csv=str(tmp_path / "a.csv"), output_json=str(tmp_path / "a.json"))
        assert out.startswith("Error: File not found"), out
        assert c.calls == []


BILLING_CSV = (
    "ip,billing_name,price_monthly\n"
    "192.0.2.10,srv-aq9001,20\n"          # exact address, already costed
    "192.0.2.11,srv-aq9002,25\n"          # exact address, no cost yet
    "198.51.100.99,srv-bv9001,30\n"       # name known, address moved
    "198.51.100.77,new box,15\n"          # only the subnet is known
    "203.0.113.5,srv-tf9001,12\n"         # a named box nobody monitors
    "203.0.113.6,,9\n"                    # nameless line item
    "203.0.113.7,srv-gs9001,0\n"          # free line: dropped by the loader
    "127.0.0.1,loop,5\n"                  # reserved: dropped by the loader
)


class TestReconcileBillingAudit:
    def _client(self):
        return RecordingClient({"host.get": HOSTS,
                                "usermacro.get": [_macro("1", "20"), _macro("3", "0")]})

    def test_every_row_lands_in_one_bucket(self, tmp_path):
        src = tmp_path / "billing.csv"
        src.write_text(BILLING_CSV)
        outdir = tmp_path / "buckets"
        out = run_tool(costs_audit, "reconcile_billing_audit", self._client(),
                       file_path=str(src), output_dir=str(outdir))
        assert out.startswith("**Reconciliation: 6 billing rows**"), out
        rows = [ln for ln in out.splitlines() if ln.startswith("| ") and "Bucket" not in ln and "---" not in ln]
        assert rows == [
            "| importable | 1 | $25.00 | Safe to import (IP match, no cost) |",
            "| already_costed | 1 | $20.00 | Skip — cost already set |",
            "| stale_ip | 1 | $30.00 | Billing team: update IP |",
            "| subnet_match | 1 | $15.00 | Likely new member — review |",
            "| onboard | 1 | $12.00 | Ops: add host to Zabbix |",
            "| cancel | 1 | $9.00 | Finance: cancel or investigate |",
        ]
        assert out.endswith(f"Wrote 6 bucket CSVs to `{outdir}`")
        assert sorted(p.name for p in outdir.iterdir()) == [
            f"billing__{b}.csv" for b in sorted(
                ["importable", "already_costed", "stale_ip", "subnet_match", "onboard", "cancel"])
        ]
        with open(outdir / "billing__stale_ip.csv", newline="") as f:
            stale = list(csv.DictReader(f))
        assert stale == [{"ip": "198.51.100.99", "name": "srv-bv9001", "price": "30.0",
                          "zabbix_host": "srv-bv9001", "zabbix_ip": "198.51.100.20"}]
        with open(outdir / "billing__subnet_match.csv", newline="") as f:
            assert list(csv.DictReader(f))[0]["zabbix_host"] == "srv-bv9001"

    def test_default_output_dir_is_next_to_the_input(self, tmp_path):
        src = tmp_path / "billing.csv"
        src.write_text(BILLING_CSV)
        out = run_tool(costs_audit, "reconcile_billing_audit", self._client(), file_path=str(src))
        assert out.endswith(f"Wrote 6 bucket CSVs to `{tmp_path}`"), out
        assert (tmp_path / "billing__importable.csv").exists()

    def test_empty_missing_and_unwritable(self, tmp_path):
        empty = tmp_path / "empty.csv"
        empty.write_text("ip,billing_name,price_monthly\n")
        c = self._client()
        assert run_tool(costs_audit, "reconcile_billing_audit", c,
                        file_path=str(empty)) == "No valid billing rows found."
        assert run_tool(costs_audit, "reconcile_billing_audit", c,
                        file_path=str(tmp_path / "nope.csv")).startswith("Failed to read CSV:")
        assert c.calls == []
        src = tmp_path / "billing.csv"
        src.write_text(BILLING_CSV)
        assert run_tool(costs_audit, "reconcile_billing_audit", c, file_path=str(src),
                        output_dir="/nonexistent-root").startswith("Failed to write bucket CSVs:")


class TestFindStaleBillingIps:
    CSV = (
        "ip,billing_name,price_monthly\n"
        "192.0.2.10,srv-aq9001,20\n"            # address still current
        "198.51.100.99,srv-bv9001,30\n"         # moved
        "203.0.113.5,srv-hm9001 extra,12\n"     # first token names the host: moved
        "203.0.113.6,unknown-um9001,7\n"        # nobody
        "203.0.113.7,,4\n"                      # no name to match on
    )

    def test_moved_addresses_are_listed_by_value(self, tmp_path):
        src = tmp_path / "billing.csv"
        src.write_text(self.CSV)
        out = run_tool(costs_audit, "find_stale_billing_ips", RecordingClient({"host.get": HOSTS}),
                       file_path=str(src))
        assert out.startswith("**2 stale billing IPs** ($42.00/mo affected)"), out
        rows = [ln for ln in out.splitlines() if ln.startswith("| srv-")]
        assert rows == [
            "| srv-bv9001 | 198.51.100.99 | 198.51.100.20 | $30.00 |",
            "| srv-hm9001 extra | 203.0.113.5 | 198.51.100.30 | $12.00 |",
        ]

    def test_nothing_stale_and_nothing_to_read(self, tmp_path):
        src = tmp_path / "billing.csv"
        src.write_text("ip,billing_name,price_monthly\n192.0.2.10,srv-aq9001,20\n")
        assert run_tool(costs_audit, "find_stale_billing_ips", RecordingClient({"host.get": HOSTS}),
                        file_path=str(src)) == "No stale billing IPs detected."
        src.write_text("ip,billing_name,price_monthly\n")
        assert run_tool(costs_audit, "find_stale_billing_ips", RecordingClient({"host.get": HOSTS}),
                        file_path=str(src)) == "No valid billing rows."
